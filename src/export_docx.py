import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from urllib.parse import urljoin
from io import BytesIO
import re
import base64
from Crypto.Cipher import AES


# ============ 0. 搜狐图片解密函数 ============
def decrypt_sohu_img(data):
    """解密搜狐文章的加密图片地址"""
    try:
        key = 'www.sohu.com6666'.encode('utf-8')
        cipher = AES.new(key, AES.MODE_ECB)
        encrypted = base64.b64decode(data)
        decrypted = cipher.decrypt(encrypted)
        # 去除PKCS7填充
        pad_len = decrypted[-1]
        if isinstance(pad_len, int) and 0 < pad_len <= 16:
            decrypted = decrypted[:-pad_len]
        result = decrypted.decode('utf-8', errors='ignore')
        # 补全协议
        if result.startswith('//'):
            result = 'https:' + result
        elif result.startswith('/'):
            result = 'https://www.sohu.com' + result
        return result
    except Exception as e:
        print(f"  解密失败：{e}")
        return None


# ============ 1. 请求网页 ============
url = "https://www.sohu.com/a/450841061_500680"
headers = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                  "AppleWebKit/537.36 (KHTML, like Gecko) "
                  "Chrome/120.0.0.0 Safari/537.36",
    "Referer": "https://www.sohu.com/"
}
print("正在请求网页...")
resp = requests.get(url, headers=headers, timeout=15)
resp.encoding = "utf-8"
print(f"状态码：{resp.status_code}")

# ============ 2. 解析页面 ============
soup = BeautifulSoup(resp.text, "html.parser")

title_tag = soup.find("h1") or soup.find("title")
article_title = title_tag.get_text(strip=True) if title_tag else "诗词题库"

content_div = (
        soup.find("article")
        or soup.find("div", class_=re.compile(r"article[_-]?content"))
        or soup.find("div", id="mp-editor")
        or soup.find("div", class_=re.compile(r"article"))
)
if content_div is None:
    all_p = soup.find_all("p")
    content_div = all_p[0].parent if all_p else None

if content_div is None:
    print("未能定位正文")
    exit(1)

# ============ 3. 创建Word文档 ============
doc = Document()


def set_cn_font(run, font_name="宋体", size=None, bold=False, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_para.add_run(article_title)
set_cn_font(run, "黑体", size=20, bold=True, color=(0x1F, 0x3A, 0x5F))


# ============ 4. 图片处理函数 ============
def add_image(img_url):
    """下载图片并插入Word"""
    if not img_url:
        return
    try:
        img_resp = requests.get(img_url, headers=headers, timeout=10)
        if img_resp.status_code == 200:
            image_stream = BytesIO(img_resp.content)
            doc.add_picture(image_stream, width=Inches(4.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            print(f"  已插入图片：{img_url}")
        else:
            print(f"  图片下载失败（状态码{img_resp.status_code}）")
    except Exception as e:
        print(f"  图片处理异常：{e}")


def get_img_url(img_tag):
    """从img标签中获取真实图片地址（处理加密data-src）"""
    # 优先检查 data-src（加密地址）
    data_src = img_tag.get('data-src')
    if data_src:
        # 判断是否为加密内容（Base64密文通常较长且无URL特征）
        if not data_src.startswith(('http', '//')):
            decrypted = decrypt_sohu_img(data_src)
            if decrypted:
                return decrypted
            # 解密失败则退回使用src
        else:
            return urljoin(url, data_src)

    # 退回普通 src
    src = img_tag.get('src')
    if src and not src.startswith('data:'):
        return urljoin(url, src)
    return None


# ============ 5. 遍历正文 ============
current_question = None

for child in content_div.children:
    if child.name == 'p':
        # 先处理 p 内的图片
        for img in child.find_all('img'):
            img_url = get_img_url(img)
            if img_url:
                add_image(img_url)

        # 处理文本
        text = child.get_text(strip=True)
        if text:
            match = re.match(r'^(\d+)[、.]\s*(.*)', text)
            if match:
                num = match.group(1)
                content = match.group(2)
                ans_split = re.split(r'答案[：:]\s*', content, maxsplit=1)
                if len(ans_split) == 2:
                    q_text = ans_split[0].strip()
                    a_text = ans_split[1].strip()
                else:
                    q_text = content.strip()
                    a_text = ""
                q_para = doc.add_paragraph()
                q_para.paragraph_format.space_before = Pt(6)
                q_para.paragraph_format.space_after = Pt(2)
                run_num = q_para.add_run(f"{num}、")
                set_cn_font(run_num, "宋体", size=12, bold=True, color=(0x2E, 0x74, 0xB5))
                run_q = q_para.add_run(q_text)
                set_cn_font(run_q, "宋体", size=12)
                if a_text:
                    a_para = doc.add_paragraph()
                    a_para.paragraph_format.left_indent = Pt(24)
                    run_a = a_para.add_run(f"答案：{a_text}")
                    set_cn_font(run_a, "楷体", size=11, color=(0xC0, 0x39, 0x2B))
                current_question = (num, q_text, a_text)
            else:
                if re.match(r'^[一二三四五六七八九十]+、', text):
                    heading_para = doc.add_paragraph()
                    heading_para.paragraph_format.space_before = Pt(12)
                    heading_para.paragraph_format.space_after = Pt(6)
                    run_h = heading_para.add_run(text)
                    set_cn_font(run_h, "黑体", size=14, bold=True, color=(0x1F, 0x3A, 0x5F))
                else:
                    ans_match = re.match(r'^答案[：:]\s*(.*)', text)
                    if ans_match and current_question:
                        a_text = ans_match.group(1).strip()
                        a_para = doc.add_paragraph()
                        a_para.paragraph_format.left_indent = Pt(24)
                        run_a = a_para.add_run(f"答案：{a_text}")
                        set_cn_font(run_a, "楷体", size=11, color=(0xC0, 0x39, 0x2B))
                    else:
                        p = doc.add_paragraph()
                        run_p = p.add_run(text)
                        set_cn_font(run_p, "宋体", size=12)

    elif child.name == 'img':
        img_url = get_img_url(child)
        if img_url:
            add_image(img_url)

# ============ 6. 保存 ============
output_file = "诗词大会题库.docx"
doc.save(output_file)
print(f"\n已保存为：{output_file}")